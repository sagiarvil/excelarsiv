import os
import glob
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def safe_unmerge_in_range(ws, min_col, min_row, max_col, max_row):
    to_unmerge = []
    for rng in list(ws.merged_cells.ranges):
        if (rng.min_col <= max_col and rng.max_col >= min_col and
            rng.min_row <= max_row and rng.max_row >= min_row):
            to_unmerge.append(rng)
    for rng in to_unmerge:
        ws.unmerge_cells(str(rng))

def safe_set(ws, coord, val, font=None, fill=None, alignment=None, border=None, num_format=None):
    cell = ws[coord]
    if type(cell).__name__ == 'MergedCell':
        return
    cell.value = val
    if font: cell.font = font
    if fill: cell.fill = fill
    if alignment: cell.alignment = alignment
    if border: cell.border = border
    if num_format: cell.number_format = num_format

def apply_crm_erp_cockpit(wb_path):
    try:
        wb = openpyxl.load_workbook(wb_path)
    except Exception as e:
        return False

    c_slate_dark = "0F172A"
    c_slate_muted = "64748B"
    c_slate_border = "E2E8F0"
    c_slate_bg = "F8FAFC"
    c_white = "FFFFFF"
    
    c_navy = "1E3A8A"
    c_navy_light = "EFF6FF"
    c_crimson = "DC2626"
    c_crimson_light = "FEF2F2"
    c_crimson_border = "FECACA"
    c_crimson_btn = "FCA5A5"
    c_emerald = "059669"
    c_emerald_dark = "065F46"
    c_emerald_light = "DCFCE7"
    c_cyan = "0284C7"
    c_cyan_light = "F1F5F9"
    c_amber_light = "FEF3C7"
    c_amber_text = "92400E"
    c_teal_light = "CCFBF1"
    c_teal_text = "115E59"

    font_title = Font(name="Segoe UI", size=15, bold=True, color=c_slate_dark)
    font_sub = Font(name="Segoe UI", size=9, italic=False, color=c_slate_muted)
    font_kpi_lbl = Font(name="Segoe UI", size=8.5, bold=True, color=c_slate_muted)
    font_kpi_val = Font(name="Segoe UI", size=20, bold=True, color=c_slate_dark)
    font_kpi_val_crimson = Font(name="Segoe UI", size=20, bold=True, color=c_crimson)
    font_kpi_sub = Font(name="Segoe UI", size=8, color=c_slate_muted)
    
    fill_white = PatternFill(start_color=c_white, end_color=c_white, fill_type="solid")
    fill_canvas = PatternFill(start_color=c_slate_bg, end_color=c_slate_bg, fill_type="solid")
    fill_dark_header = PatternFill(start_color=c_slate_dark, end_color=c_slate_dark, fill_type="solid")
    fill_emerald_action = PatternFill(start_color=c_emerald_dark, end_color=c_emerald_dark, fill_type="solid")
    
    border_card = Border(left=Side(style='thin', color=c_slate_border),
                         right=Side(style='thin', color=c_slate_border),
                         top=Side(style='thin', color=c_slate_border),
                         bottom=Side(style='thin', color=c_slate_border))

    # 1. PANO SAYFASI
    if 'PANO' in wb.sheetnames:
        ws = wb['PANO']
        if ws.views and ws.views.sheetView:
            ws.views.sheetView[0].showGridLines = False

        # Çakışan merge alanlarını temizle
        safe_unmerge_in_range(ws, 2, 2, 13, 3)
        safe_unmerge_in_range(ws, 2, 11, 7, 12)

        col_widths = {
            'A': 12, 'B': 17, 'C': 12, 'D': 11,
            'E': 17, 'F': 12, 'G': 11,
            'H': 17, 'I': 13, 'J': 11,
            'K': 17, 'L': 14, 'M': 12, 'N': 6
        }
        for col_l, w in col_widths.items():
            ws.column_dimensions[col_l].width = w

        ws.row_dimensions[1].height = 24
        ws.row_dimensions[2].height = 28
        ws.row_dimensions[3].height = 18
        ws.row_dimensions[4].height = 10
        ws.row_dimensions[5].height = 20
        ws.row_dimensions[6].height = 6
        ws.row_dimensions[7].height = 32
        ws.row_dimensions[8].height = 6
        ws.row_dimensions[9].height = 18
        ws.row_dimensions[10].height = 14
        ws.row_dimensions[11].height = 24
        ws.row_dimensions[12].height = 18
        ws.row_dimensions[13].height = 8
        ws.row_dimensions[14].height = 24
        for r in range(15, 19):
            ws.row_dimensions[r].height = 22
        ws.row_dimensions[19].height = 26
        ws.row_dimensions[20].height = 22
        ws.row_dimensions[21].height = 14
        ws.row_dimensions[22].height = 24
        for r in range(23, 27):
            ws.row_dimensions[r].height = 20
        ws.row_dimensions[27].height = 26

        safe_set(ws, 'A1', "☰ PANO", font=Font(name="Segoe UI", size=9, bold=True, color=c_slate_dark),
                 fill=fill_white, alignment=Alignment(horizontal='center', vertical='center'), border=border_card)

        curr_title = ws['B2'].value or "STRATEJİK YÖNETİCİ KOKPİTİ"
        try:
            ws.merge_cells('B2:K2')
        except Exception:
            pass
        safe_set(ws, 'B2', str(curr_title).upper(), font=font_title, alignment=Alignment(vertical='center'))

        safe_set(ws, 'L2', "Dönem: 2026 / Tam ▼", font=Font(name="Segoe UI", size=8.5, color="334155"),
                 fill=fill_white, alignment=Alignment(horizontal='center', vertical='center'), border=border_card)

        safe_set(ws, 'M2', "★ DOĞRULANMIŞ", font=Font(name="Segoe UI", size=8.5, bold=True, color=c_white),
                 fill=PatternFill(start_color=c_emerald, end_color=c_emerald, fill_type="solid"),
                 alignment=Alignment(horizontal='center', vertical='center'))

        curr_sub = ws['B3'].value or "Yönetici Karar Kokpiti, Otomatik Doğrulama ve Dinamik Hesaplama Sistemi · 2026"
        try:
            ws.merge_cells('B3:K3')
        except Exception:
            pass
        safe_set(ws, 'B3', str(curr_sub), font=font_sub, alignment=Alignment(vertical='center'))

        kpi_configs = [
            {'start_col': 'B', 'end_col': 'D', 'accent': c_navy, 'badge_fill': c_navy_light, 'badge_color': "1E40AF"},
            {'start_col': 'E', 'end_col': 'G', 'accent': c_crimson, 'badge_fill': c_crimson_light, 'badge_color': "B91C1C"},
            {'start_col': 'H', 'end_col': 'J', 'accent': c_emerald, 'badge_fill': c_emerald_light, 'badge_color': "15803D"},
            {'start_col': 'K', 'end_col': 'M', 'accent': c_cyan, 'badge_fill': c_cyan_light, 'badge_color': "475569"}
        ]

        for cfg in kpi_configs:
            sc = openpyxl.utils.column_index_from_string(cfg['start_col'])
            ec = openpyxl.utils.column_index_from_string(cfg['end_col'])
            
            for r in range(5, 10):
                for c in range(sc, ec + 1):
                    cell = ws.cell(row=r, column=c)
                    if type(cell).__name__ == 'MergedCell': continue
                    cell.fill = fill_white
                    left_s = Side(style='medium', color=cfg['accent']) if c == sc else Side(style='none')
                    right_s = Side(style='thin', color=c_slate_border) if c == ec else Side(style='none')
                    top_s = Side(style='thin', color=c_slate_border) if r == 5 else Side(style='none')
                    bottom_s = Side(style='thin', color=c_slate_border) if r == 9 else Side(style='none')
                    cell.border = Border(left=left_s, right=right_s, top=top_s, bottom=bottom_s)

            c_lbl = ws.cell(row=5, column=sc)
            if type(c_lbl).__name__ != 'MergedCell':
                c_lbl.font = font_kpi_lbl
                c_lbl.alignment = Alignment(vertical='center', indent=1)

            c_bdg = ws.cell(row=5, column=ec)
            if type(c_bdg).__name__ != 'MergedCell':
                c_bdg.font = Font(name="Segoe UI", size=7.5, bold=True, color=cfg['badge_color'])
                c_bdg.fill = PatternFill(start_color=cfg['badge_fill'], end_color=cfg['badge_fill'], fill_type="solid")
                c_bdg.alignment = Alignment(horizontal='center', vertical='center')
                c_bdg.border = border_card

            c_val = ws.cell(row=7, column=sc)
            if type(c_val).__name__ != 'MergedCell':
                c_val.font = font_kpi_val_crimson if cfg['accent'] == c_crimson else font_kpi_val
                c_val.alignment = Alignment(vertical='center', indent=1)

            c_sub = ws.cell(row=9, column=sc)
            if type(c_sub).__name__ != 'MergedCell':
                c_sub.font = font_kpi_sub
                c_sub.alignment = Alignment(vertical='center', indent=1)

        safe_set(ws, 'B11', ws['B11'].value or "DÖNEMSEL İCRA VE PROJEKSİYON",
                 font=Font(name="Segoe UI", size=10.5, bold=True, color=c_slate_dark))
        safe_set(ws, 'B12', ws['B12'].value or "Kümülatif Hacim vs. Operasyonel Trend", font=font_sub)
        safe_set(ws, 'F11', "+%32 Verim", font=Font(name="Segoe UI", size=8, bold=True, color="15803D"),
                 fill=PatternFill(start_color=c_emerald_light, end_color=c_emerald_light, fill_type="solid"),
                 alignment=Alignment(horizontal='center', vertical='center'), border=border_card)

        for col_l in ['B', 'C', 'D', 'E']:
            cell = ws[f'{col_l}14']
            if type(cell).__name__ == 'MergedCell': continue
            cell.font = Font(name="Segoe UI", size=8.5, bold=True, color="475569")
            cell.fill = PatternFill(start_color=c_slate_bg, end_color=c_slate_bg, fill_type="solid")
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = border_card
        
        for r in range(15, 21):
            bg_c = c_white if r % 2 == 1 else c_slate_bg
            f_fill = PatternFill(start_color=bg_c, end_color=bg_c, fill_type="solid")
            for col_l in ['B', 'C', 'D', 'E']:
                cell = ws[f'{col_l}{r}']
                if type(cell).__name__ == 'MergedCell': continue
                cell.font = Font(name="Segoe UI", size=9, color=c_slate_dark)
                cell.fill = f_fill
                cell.border = border_card
                cell.alignment = Alignment(horizontal='center', vertical='center')

        safe_set(ws, 'H11', ws['H11'].value or "KONSOLİDE OPERASYON VE KARAR MATRİSİ",
                 font=Font(name="Segoe UI", size=10.5, bold=True, color=c_slate_dark))
        safe_set(ws, 'H12', ws['H12'].value or "Anlık Hesaplama Sonuçları ve Durum Puanlaması", font=font_sub)
        
        for c_idx in range(8, 14):
            cell = ws.cell(row=14, column=c_idx)
            if type(cell).__name__ == 'MergedCell': continue
            cell.fill = fill_dark_header
            cell.font = Font(name="Segoe UI", size=8.5, bold=True, color=c_white)
            cell.border = Border(top=Side(style='thin', color=c_slate_dark),
                                 bottom=Side(style='thin', color=c_slate_dark),
                                 left=Side(style='thin', color="334155"),
                                 right=Side(style='thin', color="334155"))
            cell.alignment = Alignment(horizontal='center' if c_idx >= 9 else 'left', vertical='center', indent=1 if c_idx==8 else 0)

        pill_styles = {
            15: {'bg': c_emerald_light, 'color': "166534"},
            16: {'bg': c_navy_light, 'color': "1E40AF"},
            17: {'bg': c_amber_light, 'color': c_amber_text},
            18: {'bg': c_teal_light, 'color': c_teal_text}
        }

        for r in range(15, 19):
            bg_c = c_white if r % 2 == 1 else c_slate_bg
            r_fill = PatternFill(start_color=bg_c, end_color=bg_c, fill_type="solid")
            for c_idx in range(8, 14):
                cell = ws.cell(row=r, column=c_idx)
                if type(cell).__name__ == 'MergedCell': continue
                cell.fill = r_fill
                cell.border = border_card
                if c_idx == 8:
                    cell.font = Font(name="Segoe UI", size=9, bold=True, color=c_slate_dark)
                    cell.alignment = Alignment(vertical='center', indent=1)
                elif c_idx in (9, 10):
                    cell.font = Font(name="Segoe UI", size=9, bold=True, color=c_slate_dark if c_idx==9 else c_crimson)
                    cell.alignment = Alignment(horizontal='right', vertical='center')
                    cell.number_format = '₺#,##0'
                elif c_idx == 11:
                    cell.font = Font(name="Segoe UI", size=8.5, color="475569")
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                elif c_idx in (12, 13):
                    if c_idx == 12:
                        p_info = pill_styles[r]
                        cell.fill = PatternFill(start_color=p_info['bg'], end_color=p_info['bg'], fill_type="solid")
                        cell.font = Font(name="Segoe UI", size=8, bold=True, color=p_info['color'])
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                        cell.border = border_card

        for c_idx in range(8, 14):
            cell = ws.cell(row=19, column=c_idx)
            if type(cell).__name__ == 'MergedCell': continue
            cell.fill = fill_dark_header
            cell.border = Border(top=Side(style='medium', color="334155"),
                                 bottom=Side(style='medium', color=c_slate_dark),
                                 left=Side(style='thin', color="334155"),
                                 right=Side(style='thin', color="334155"))
            if c_idx == 8:
                cell.font = Font(name="Segoe UI", size=9, bold=True, color=c_white)
                cell.alignment = Alignment(vertical='center', indent=1)
            elif c_idx == 9:
                cell.font = Font(name="Segoe UI", size=9.5, bold=True, color="38BDF8")
                cell.alignment = Alignment(horizontal='right', vertical='center')
                cell.number_format = '₺#,##0'
            elif c_idx == 10:
                cell.font = Font(name="Segoe UI", size=9.5, bold=True, color="F87171")
                cell.alignment = Alignment(horizontal='right', vertical='center')
                cell.number_format = '₺#,##0'
            elif c_idx in (12, 13):
                cell.font = Font(name="Segoe UI", size=8, bold=True, color="34D399")
                cell.alignment = Alignment(horizontal='center', vertical='center')

        for r in range(22, 28):
            for c in range(2, 6):
                cell = ws.cell(row=r, column=c)
                if type(cell).__name__ == 'MergedCell': continue
                cell.fill = fill_white
                top_s = Side(style='thin', color=c_slate_border) if r == 22 else Side(style='none')
                bot_s = Side(style='thin', color=c_slate_border) if r == 27 else Side(style='none')
                left_s = Side(style='thin', color=c_slate_border) if c == 2 else Side(style='none')
                right_s = Side(style='thin', color=c_slate_border) if c == 5 else Side(style='none')
                cell.border = Border(top=top_s, bottom=bot_s, left=left_s, right=right_s)

        safe_set(ws, 'B22', ws['B22'].value or "KONSOLİDE PAY VE AĞIRLIK",
                 font=Font(name="Segoe UI", size=9.5, bold=True, color=c_slate_dark),
                 alignment=Alignment(vertical='center', indent=1))
        
        segment_colors = ['059669', '0284C7', 'D97706', '8B5CF6']
        for idx, r in enumerate(range(23, 27)):
            cell = ws.cell(row=r, column=2)
            if type(cell).__name__ == 'MergedCell': continue
            cell.font = Font(name="Segoe UI", size=8.5, bold=True, color=segment_colors[idx])
            cell.alignment = Alignment(vertical='center', indent=1)

        for r in range(22, 28):
            for c in range(6, 10):
                cell = ws.cell(row=r, column=c)
                if type(cell).__name__ == 'MergedCell': continue
                cell.fill = fill_white
                top_s = Side(style='thin', color=c_slate_border) if r == 22 else Side(style='none')
                bot_s = Side(style='thin', color=c_slate_border) if r == 27 else Side(style='none')
                left_s = Side(style='thin', color=c_slate_border) if c == 6 else Side(style='none')
                right_s = Side(style='thin', color=c_slate_border) if c == 9 else Side(style='none')
                cell.border = Border(top=top_s, bottom=bot_s, left=left_s, right=right_s)

        safe_set(ws, 'F22', ws['F22'].value or "LİKİDİTE RİSKİ & ERKEN UYARI SİNYALLERİ",
                 font=Font(name="Segoe UI", size=9.5, bold=True, color=c_slate_dark),
                 alignment=Alignment(vertical='center', indent=1))
        
        safe_set(ws, 'F23', ws['F23'].value or "Nakit Açığı Riski (30 Gün)", font=Font(name="Segoe UI", size=8.5, color="475569"))
        safe_set(ws, 'I23', ws['I23'].value or "0.00 TL (Sıfır Risk)", font=Font(name="Segoe UI", size=8.5, bold=True, color=c_emerald))
        safe_set(ws, 'F24', ws['F24'].value or "Tahsilat Gecikme Toleransı", font=Font(name="Segoe UI", size=8.5, color="475569"))
        safe_set(ws, 'I24', ws['I24'].value or "48 Gün Güvenli Eşik", font=Font(name="Segoe UI", size=8.5, bold=True, color=c_emerald))

        for c in range(6, 10):
            cell = ws.cell(row=25, column=c)
            if type(cell).__name__ == 'MergedCell': continue
            cell.fill = PatternFill(start_color=c_emerald_light, end_color=c_emerald_light, fill_type="solid")
            cell.border = border_card
        safe_set(ws, 'F25', ws['F25'].value or "██████████████████████████████ Güvenli Bölge (%100)",
                 font=Font(name="Segoe UI", size=8, bold=True, color="166534"),
                 alignment=Alignment(horizontal='center', vertical='center'))

        for c in range(6, 10):
            cell = ws.cell(row=26, column=c)
            if type(cell).__name__ == 'MergedCell': continue
            cell.fill = PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid")
            cell.border = border_card
        safe_set(ws, 'F26', ws['F26'].value or "✓ 90 GÜNLÜK PROJEKSİYONDA NAKİT SIKIŞIKLIĞI TESPİT EDİLMEDİ",
                 font=Font(name="Segoe UI", size=8, bold=True, color=c_emerald),
                 alignment=Alignment(horizontal='center', vertical='center'))

        for r in range(22, 28):
            for c in range(11, 14):
                cell = ws.cell(row=r, column=c)
                if type(cell).__name__ == 'MergedCell': continue
                cell.fill = fill_emerald_action
                top_s = Side(style='thin', color=c_emerald_dark) if r == 22 else Side(style='none')
                bot_s = Side(style='thin', color=c_emerald_dark) if r == 27 else Side(style='none')
                left_s = Side(style='thin', color=c_emerald_dark) if c == 11 else Side(style='none')
                right_s = Side(style='thin', color=c_emerald_dark) if c == 13 else Side(style='none')
                cell.border = Border(top=top_s, bottom=bot_s, left=left_s, right=right_s)

        safe_set(ws, 'K22', "KARAR MOTORU", font=Font(name="Segoe UI", size=7.5, bold=True, color="A7F3D0"),
                 alignment=Alignment(vertical='center', indent=1))
        safe_set(ws, 'K23', "LİKİDİTE OPTİMİZASYON FIRSATI", font=Font(name="Segoe UI", size=10, bold=True, color=c_white),
                 alignment=Alignment(vertical='center', indent=1))
        safe_set(ws, 'K24', "Atıl kalan nakit rezervinin gecelik repoda", font=Font(name="Segoe UI", size=8, color="D1FAE5"),
                 alignment=Alignment(vertical='center', indent=1))
        safe_set(ws, 'K25', "değerlendirilmesiyle ek getiri imkanı.", font=Font(name="Segoe UI", size=8, color="D1FAE5"),
                 alignment=Alignment(vertical='center', indent=1))

        for c in range(11, 14):
            cell = ws.cell(row=26, column=c)
            if type(cell).__name__ == 'MergedCell': continue
            cell.fill = fill_white
            cell.border = border_card
        safe_set(ws, 'K26', "FONLAMA SENARYOSUNU İNCELE ›", font=Font(name="Segoe UI", size=8.5, bold=True, color=c_emerald_dark),
                 alignment=Alignment(horizontal='center', vertical='center'))

    # 2. YÖNETİM_RAPORU SAYFASI
    if 'YÖNETİM_RAPORU' in wb.sheetnames:
        ws_r = wb['YÖNETİM_RAPORU']
        if ws_r.views and ws_r.views.sheetView:
            ws_r.views.sheetView[0].showGridLines = False

        for col_l, w in [('A', 12), ('B', 24), ('C', 14), ('D', 14), ('E', 16), ('F', 14),
                         ('G', 6),  ('H', 24), ('I', 14), ('J', 14), ('K', 16), ('L', 14), ('M', 6)]:
            ws_r.column_dimensions[col_l].width = w

        ws_r.row_dimensions[1].height = 24
        ws_r.row_dimensions[2].height = 28
        ws_r.row_dimensions[3].height = 18
        ws_r.row_dimensions[4].height = 20
        ws_r.row_dimensions[5].height = 12
        ws_r.row_dimensions[6].height = 26
        for r in range(7, 14):
            ws_r.row_dimensions[r].height = 22
        ws_r.row_dimensions[14].height = 28
        ws_r.row_dimensions[15].height = 14
        ws_r.row_dimensions[16].height = 10
        ws_r.row_dimensions[17].height = 26
        for r in range(18, 28):
            ws_r.row_dimensions[r].height = 22

        safe_set(ws_r, 'A1', "☰ PANO", font=Font(name="Segoe UI", size=9, bold=True, color=c_slate_dark),
                 fill=fill_white, alignment=Alignment(horizontal='center', vertical='center'), border=border_card)

        safe_set(ws_r, 'B2', ws_r['B2'].value or "NİHAİ KARAR VE DEĞERLEME RAPORU",
                 font=Font(name="Segoe UI", size=14, bold=True, color=c_slate_dark))
        safe_set(ws_r, 'B3', ws_r['B3'].value or "Yönetim Kurulu ve İcra Komitesi için Otomatik Finansal Stres Testi Raporu", font=font_sub)
        safe_set(ws_r, 'B4', "DÖNEM: 2026 Q3 · DENETİM PROTOKOLÜ: ISO/VUK/IFRS UYUMLU · DURUM: ONAYLANDI",
                 font=Font(name="Segoe UI", size=8.5, bold=True, color=c_emerald))

        for r in range(6, 15):
            for c in range(2, 7):
                cell = ws_r.cell(row=r, column=c)
                if type(cell).__name__ == 'MergedCell': continue
                if r == 6:
                    cell.fill = PatternFill(start_color=c_crimson_light, end_color=c_crimson_light, fill_type="solid")
                    cell.border = Border(top=Side(style='thin', color=c_crimson_border),
                                         bottom=Side(style='thin', color=c_crimson_border),
                                         left=Side(style='thin', color=c_crimson_border) if c==2 else Side(style='none'),
                                         right=Side(style='thin', color=c_crimson_border) if c==6 else Side(style='none'))
                elif r == 14:
                    cell.fill = PatternFill(start_color=c_crimson_light, end_color=c_crimson_light, fill_type="solid")
                    cell.border = Border(top=Side(style='thin', color=c_crimson_btn),
                                         bottom=Side(style='thin', color=c_crimson_btn),
                                         left=Side(style='thin', color=c_crimson_btn) if c==2 else Side(style='none'),
                                         right=Side(style='thin', color=c_crimson_btn) if c==6 else Side(style='none'))
                else:
                    cell.fill = fill_white
                    cell.border = Border(left=Side(style='thin', color=c_crimson_border) if c==2 else Side(style='none'),
                                         right=Side(style='thin', color=c_crimson_border) if c==6 else Side(style='none'),
                                         top=Side(style='none'), bottom=Side(style='none'))

        safe_set(ws_r, 'B6', "SENARYO A: KÖTÜMSER (STRES TESTİ -%20)",
                 font=Font(name="Segoe UI", size=9.5, bold=True, color="991B1B"),
                 alignment=Alignment(vertical='center', indent=1))

        safe_set(ws_r, 'B8', "Ciro / Hacim Düşüşü:", font=Font(name="Segoe UI", size=9, color="475569"))
        safe_set(ws_r, 'E8', "-%20.0 (Kriz Eşiği)", font=Font(name="Segoe UI", size=9.5, bold=True, color=c_crimson),
                 alignment=Alignment(horizontal='right', vertical='center'))

        safe_set(ws_r, 'B9', "Beklenen Net Kâr / Getiri:", font=Font(name="Segoe UI", size=9, color="475569"))
        safe_set(ws_r, 'E9', ws_r['E9'].value or 4280000, font=Font(name="Segoe UI", size=9.5, bold=True, color=c_slate_dark),
                 alignment=Alignment(horizontal='right', vertical='center'), num_format='₺#,##0')

        safe_set(ws_r, 'B10', "Minimum Nakit Tamponu:", font=Font(name="Segoe UI", size=9, color="475569"))
        safe_set(ws_r, 'E10', ws_r['E10'].value or "₺1.850.000 (Yeterli)", font=Font(name="Segoe UI", size=9.5, bold=True, color=c_emerald),
                 alignment=Alignment(horizontal='right', vertical='center'))

        safe_set(ws_r, 'B11', "Likidite Karşılama Oranı:", font=Font(name="Segoe UI", size=9, color="475569"))
        safe_set(ws_r, 'E11', ws_r['E11'].value or "1.28x (Eşik: 1.10x)", font=Font(name="Segoe UI", size=9.5, bold=True, color=c_emerald),
                 alignment=Alignment(horizontal='right', vertical='center'))

        safe_set(ws_r, 'B14', "STRES SENARYOSUNDA DAHİ İFLAS RİSKİ YOK", font=Font(name="Segoe UI", size=8.5, bold=True, color="B91C1C"),
                 alignment=Alignment(horizontal='center', vertical='center'))

        for r in range(6, 15):
            for c in range(8, 13):
                cell = ws_r.cell(row=r, column=c)
                if type(cell).__name__ == 'MergedCell': continue
                if r == 6:
                    cell.fill = PatternFill(start_color=c_navy_light, end_color=c_navy_light, fill_type="solid")
                    cell.border = Border(top=Side(style='thin', color="BFDBFE"),
                                         bottom=Side(style='thin', color="BFDBFE"),
                                         left=Side(style='thin', color="BFDBFE") if c==8 else Side(style='none'),
                                         right=Side(style='thin', color="BFDBFE") if c==12 else Side(style='none'))
                elif r == 14:
                    cell.fill = PatternFill(start_color=c_navy_light, end_color=c_navy_light, fill_type="solid")
                    cell.border = Border(top=Side(style='thin', color="93C5FD"),
                                         bottom=Side(style='thin', color="93C5FD"),
                                         left=Side(style='thin', color="93C5FD") if c==8 else Side(style='none'),
                                         right=Side(style='thin', color="93C5FD") if c==12 else Side(style='none'))
                else:
                    cell.fill = fill_white
                    cell.border = Border(left=Side(style='thin', color="BFDBFE") if c==8 else Side(style='none'),
                                         right=Side(style='thin', color="BFDBFE") if c==12 else Side(style='none'),
                                         top=Side(style='none'), bottom=Side(style='none'))

        safe_set(ws_r, 'H6', "SENARYO B: BAZ MODEL (HEDEFLENEN)",
                 font=Font(name="Segoe UI", size=9.5, bold=True, color="1E40AF"),
                 alignment=Alignment(vertical='center', indent=1))

        safe_set(ws_r, 'H8', "Ciro / Hacim Gerçekleşmesi:", font=Font(name="Segoe UI", size=9, color="475569"))
        safe_set(ws_r, 'K8', "%100.0 (Tam Bütçe)", font=Font(name="Segoe UI", size=9.5, bold=True, color="2563EB"),
                 alignment=Alignment(horizontal='right', vertical='center'))

        safe_set(ws_r, 'H9', "Beklenen Net Kâr / Getiri:", font=Font(name="Segoe UI", size=9, color="475569"))
        safe_set(ws_r, 'K9', ws_r['K9'].value or 12450000, font=Font(name="Segoe UI", size=9.5, bold=True, color=c_slate_dark),
                 alignment=Alignment(horizontal='right', vertical='center'), num_format='₺#,##0')

        safe_set(ws_r, 'H10', "Serbest Nakit Akımı:", font=Font(name="Segoe UI", size=9, color="475569"))
        safe_set(ws_r, 'K10', ws_r['K10'].value or 8240000, font=Font(name="Segoe UI", size=9.5, bold=True, color=c_emerald),
                 alignment=Alignment(horizontal='right', vertical='center'), num_format='₺#,##0')

        safe_set(ws_r, 'H11', "Yatırım Geri Dönüşü (ROI):", font=Font(name="Segoe UI", size=9, color="475569"))
        safe_set(ws_r, 'K11', ws_r['K11'].value or "%38.4 Yıllık", font=Font(name="Segoe UI", size=9.5, bold=True, color=c_emerald),
                 alignment=Alignment(horizontal='right', vertical='center'))

        safe_set(ws_r, 'H14', "OPTİMUM İCRA VE PLANLAMA REÇETESİ", font=Font(name="Segoe UI", size=8.5, bold=True, color="1D4ED8"),
                 alignment=Alignment(horizontal='center', vertical='center'))

        for r in range(17, 28):
            for c in range(2, 13):
                cell = ws_r.cell(row=r, column=c)
                if type(cell).__name__ == 'MergedCell': continue
                cell.fill = fill_white
                top_s = Side(style='thin', color=c_slate_border) if r == 17 else Side(style='none')
                bot_s = Side(style='thin', color=c_slate_border) if r == 27 else Side(style='none')
                left_s = Side(style='thin', color=c_slate_border) if c == 2 else Side(style='none')
                right_s = Side(style='thin', color=c_slate_border) if c == 12 else Side(style='none')
                cell.border = Border(top=top_s, bottom=bot_s, left=left_s, right=right_s)

        safe_set(ws_r, 'B17', "YÖNETİM KARAR MATRİSİ VE UYGULAMA REÇETESİ",
                 font=Font(name="Segoe UI", size=10.5, bold=True, color=c_slate_dark),
                 alignment=Alignment(vertical='center', indent=1))

        priority_rows = [(19, 20), (21, 22), (23, 24), (25, 26)]
        for h_row, b_row in priority_rows:
            safe_set(ws_r, f'B{h_row}', ws_r[f'B{h_row}'].value, font=Font(name="Segoe UI", size=9.5, bold=True, color=c_slate_dark),
                     alignment=Alignment(vertical='center', indent=1))
            safe_set(ws_r, f'B{b_row}', ws_r[f'B{b_row}'].value, font=Font(name="Segoe UI", size=8.5, color="475569"),
                     alignment=Alignment(vertical='center', indent=2))

    wb.save(wb_path)
    return True

if __name__ == '__main__':
    all_files = sorted(glob.glob('EXCELARSIV_PRODUCTS/**/*.xlsx', recursive=True))
    count = 0
    for idx, p in enumerate(all_files, 1):
        if apply_crm_erp_cockpit(p):
            count += 1
            if count % 20 == 0 or count == len(all_files):
                print(f"[{count}/{len(all_files)}] Processed {os.path.basename(p)}")
    print(f"Successfully applied CRM/ERP Cockpit styling to {count} product workbooks.")
