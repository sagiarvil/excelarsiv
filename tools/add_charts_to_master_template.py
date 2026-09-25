import openpyxl
from openpyxl.chart import BarChart, DoughnutChart, Reference

master_path = 'calisma_prorokolu/EXCELARSIV_DECISION_OS_V15_1_MASTER_STANDART_UI_LOCKED.xlsx'
wb = openpyxl.load_workbook(master_path)

# 03_GORSEL_TASARIM içine kural ekle
if '03_GORSEL_TASARIM' in wb.sheetnames:
    ws_v = wb['03_GORSEL_TASARIM']
    next_r = ws_v.max_row + 1
    ws_v[f'B{next_r}'] = "Vitrin & Yönetici Grafikleri (Executive Showcase Charts)"
    ws_v[f'C{next_r}'] = "PANO (Dashboard) sayfasında en az 2 adet yüksek estetikli dinamik grafik bulunmalıdır: 1) Dönemsel Performans & Trend (Çoklu Sütun Bar Grafiği), 2) Portföy / Risk Dağılımı (Modern İnce Halka - Doughnut Chart, HoleSize=%60). Grafikler Bento Grid simetrisinde B28:G44 ve H28:M44 koordinatlarına oturmalıdır."
    ws_v[f'D{next_r}'] = "ZORUNLU"

wb.save(master_path)
print("Master Template: Yönetici grafikleri kuralı başarıyla eklendi.")
