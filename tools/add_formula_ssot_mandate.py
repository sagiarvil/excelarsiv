with open('calisma_prorokolu/EXCELARSIV_DECISION_OS_V15_1_RUNTIME_KERNEL.py', 'r', encoding='utf-8') as f:
    content = f.read()

rule_257 = '''
    # MANDATE RULE 257: CANONICAL OPENXML FORMULA SSOT & ZERO #NAME? DEFECTS
    # 1. OpenXML Evrensel Formül Sözleşmesi: Excel dosya formatı (OOXML) gereği, hücre formüllerinde asla yerel
    #    Türkçe fonksiyon adları (TOPLA, EĞER, DÜŞEYARA, YUVARLA vb.) ve noktalı virgül (;) parametre ayırıcısı kullanılamaz.
    #    Tüm formüller evrensel standartta (SUM, IF, VLOOKUP, ROUND vb.) ve virgül (,) ayırıcısı ile yazılmalıdır.
    # 2. Canlı Bağlayıcı Bütünlüğü: PANO ve Dashboard KPI kartları (B7, E7, H7, K7), GİRDİLER ve MOTOR sekmelerindeki
    #    toplam ve hesaplama hücrelerine canlı ve kırılmasız formüllerle bağlanmalıdır; hiçbir hücrede #AD? (#NAME?)
    #    veya #BAŞV! (#REF!) hatası bulunamaz.
'''

if 'MANDATE RULE 257' not in content:
    target = '    # MANDATE RULE 256: ZERO LABEL-DATA OVERLAP & ISOLATED CHART ANCHORING'
    content = content.replace(target, rule_257 + '\n' + target)
    with open('calisma_prorokolu/EXCELARSIV_DECISION_OS_V15_1_RUNTIME_KERNEL.py', 'w', encoding='utf-8') as f:
        f.write(content)
    print("Rule 257 Kernel py dosyasına başarıyla işlendi.")

import openpyxl

master_path = 'calisma_prorokolu/EXCELARSIV_DECISION_OS_V15_1_MASTER_STANDART_UI_LOCKED.xlsx'
wb = openpyxl.load_workbook(master_path)

if '03_GORSEL_TASARIM' in wb.sheetnames:
    ws_v = wb['03_GORSEL_TASARIM']
    next_r = ws_v.max_row + 1
    ws_v[f'B{next_r}'] = "OpenXML Evrensel Formül Sözleşmesi & Sıfır #AD? Hatası"
    ws_v[f'C{next_r}'] = "Tüm hücre formülleri evrensel İngilizce fonksiyonlar (SUM, IF, ROUND, VLOOKUP) ve virgül (,) parametre ayracı ile yazılmalıdır. Türkçe fonksiyon adları veya noktalı virgül (;) kullanımından kaynaklanan #AD? (#NAME?) hataları kesinlikle yasaktır."
    ws_v[f'D{next_r}'] = "ZORUNLU (SSOT)"
    wb.save(master_path)
    print("Master Template: Rule 257 başarıyla işlendi.")
