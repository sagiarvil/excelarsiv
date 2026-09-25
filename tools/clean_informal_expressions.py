import openpyxl

# 1. BankaKrediVeTaksitTakipSistemi.xlsx temizliği
prod_path = 'EXCELARSIV_PRODUCTS/banka-kredi-ve-taksit-takip-sistemi/BankaKrediVeTaksitTakipSistemi.xlsx'
wb_p = openpyxl.load_workbook(prod_path)

# KILAVUZ C9
ws_k = wb_p['KILAVUZ']
ws_k['C9'] = "Her kademedeki yöneticinin anında kavrayabileceği yalınlıkta sistem özeti, kullanım kılavuzu ve finans terimleri sözlüğü."

# 00_HAKKINDA_VE_REHBER B3 ve C7
ws_info = wb_p['00_HAKKINDA_VE_REHBER']
ws_info['B3'] = "Yalın, Sade ve Eksiksiz Finansal Karar Mimarisi & Kullanım Rehberi · Siemens & Enterprise Standartları"
ws_info['C7'] = "YALIN KURUMSAL AÇIKLAMA"

wb_p.save(prod_path)
print("Ürün dosyasındaki gayriresmi ifadeler kurumsal C-Level dile uyarlandı.")

# 2. Master Şablon temizliği
master_path = 'calisma_prorokolu/EXCELARSIV_DECISION_OS_V15_1_MASTER_STANDART_UI_LOCKED.xlsx'
wb_m = openpyxl.load_workbook(master_path)

if '00_HAKKINDA_VE_REHBER' in wb_m.sheetnames:
    ws_mi = wb_m['00_HAKKINDA_VE_REHBER']
    ws_mi['B3'] = "Yalın, Sade ve Eksiksiz Finansal Karar Mimarisi & Kullanım Rehberi · Siemens & Enterprise Standartları"
    ws_mi['C7'] = "YALIN KURUMSAL AÇIKLAMA"

if '03_GORSEL_TASARIM' in wb_m.sheetnames:
    ws_mv = wb_m['03_GORSEL_TASARIM']
    for r in range(1, ws_mv.max_row + 1):
        val = str(ws_mv.cell(r, 3).value or '')
        if '5 yaş' in val or 'çocuk' in val:
            ws_mv.cell(r, 3).value = "00_HAKKINDA_VE_REHBER (en üst düzey yalınlıkta sistem sözlüğü ve rehberi), kart içi sıfır gridline/artık border, simetrik dış çerçeve ve altın oran sütun genişlikleri zorunludur."

wb_m.save(master_path)
print("Master şablondaki gayriresmi ifadeler kurumsal C-Level dile uyarlandı.")
