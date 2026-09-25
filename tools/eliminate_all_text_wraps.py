import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side

wb_path = 'EXCELARSIV_PRODUCTS/banka-kredi-ve-taksit-takip-sistemi/BankaKrediVeTaksitTakipSistemi.xlsx'
wb = openpyxl.load_workbook(wb_path)

# 1. VERİ SAYFALARINDAKİ BAŞLIK VE ALT BAŞLIKLAR (A1 ve A3)
# Bu sayfalarda A1 ve A3 tablonun genişliği kadar merge edilir, wrap_text=False yapılır.
data_sheets = [
    'BANKALAR', 'LIMITLER', 'KREDILER', 'TAKSIT_PLANI', 
    'ODEME_TAKVIMI', 'MASRAF_KOMISYON', 'TEMINATLAR', 
    'YENI_TEKLIFLER', 'REFINANSMAN', 'KARAR_MOTORU',
    'AYARLAR', 'LISTELER', 'KONTROLLER', 'SENARYO_DUYARLILIK',
    'ORNEK_VERI', 'TESTLER', 'DEGISIKLIK_KAYDI'
]

navy_fill = PatternFill(start_color='1B365D', end_color='1B365D', fill_type='solid') # Kurumsal Lacivert
white_bold = Font(name='Montserrat', size=13, bold=True, color='FFFFFF')
sub_font = Font(name='Segoe UI', size=9.5, italic=True, color='475569')

for s in data_sheets:
    if s not in wb.sheetnames:
        continue
    ws = wb[s]
    ws.views.sheetView[0].showGridLines = False
    
    # Tablonun en sağ sütununu tespit et (5. veya 4. satırdaki başlıklar)
    table_cols = [c for c in range(1, 30) if ws.cell(5, c).value is not None or ws.cell(4, c).value is not None]
    max_c = max(table_cols) if table_cols else 10
    if max_c < 8:
        max_c = 9 # Minimum genişlik garantisi
    col_let = openpyxl.utils.get_column_letter(max_c)
    
    # 1. Mevcut A1 merge'ini kaldırıp tablo genişliğinde yeniden merge et
    for mr in list(ws.merged_cells.ranges):
        if 'A1' in mr or 'A3' in mr:
            ws.merged_cells.remove(mr)
            
    # A1: Banner
    ws.merge_cells(f'A1:{col_let}1')
    ws['A1'].font = white_bold
    ws['A1'].alignment = Alignment(horizontal='left', vertical='center', wrap_text=False)
    ws.row_dimensions[1].height = 32.0
    
    # A1 banner arka planını düzgün doldur
    for c in range(1, max_c + 1):
        ws.cell(1, c).fill = navy_fill
        
    # A3: Alt Başlık (Kırılmayı sıfırlamak için tablo genişliğinde merge)
    if ws['A3'].value:
        ws.merge_cells(f'A3:{col_let}3')
        ws['A3'].font = sub_font
        ws['A3'].alignment = Alignment(horizontal='left', vertical='center', wrap_text=False)
        ws.row_dimensions[3].height = 22.0
        
    # Boşluk satırları
    ws.row_dimensions[2].height = 7.8
    ws.row_dimensions[4].height = 10.0

# 2. PANO, YÖNETİM_RAPORU, KILAVUZ ve 00_HAKKINDA_VE_REHBER sayfalarındaki başlıkları merge et
# PANO
ws_pano = wb['PANO']
for mr in list(ws_pano.merged_cells.ranges):
    if 'B2' in mr or 'B3' in mr:
        ws_pano.merged_cells.remove(mr)
ws_pano.merge_cells('B2:M2')
ws_pano['B2'].alignment = Alignment(horizontal='left', vertical='center', wrap_text=False)
ws_pano.merge_cells('B3:M3')
ws_pano['B3'].alignment = Alignment(horizontal='left', vertical='center', wrap_text=False)

# YÖNETİM_RAPORU
ws_rap = wb['YÖNETİM_RAPORU']
for mr in list(ws_rap.merged_cells.ranges):
    if 'B2' in mr or 'B3' in mr or 'B4' in mr:
        ws_rap.merged_cells.remove(mr)
ws_rap.merge_cells('B2:L2')
ws_rap['B2'].alignment = Alignment(horizontal='left', vertical='center', wrap_text=False)
ws_rap.merge_cells('B3:L3')
ws_rap['B3'].alignment = Alignment(horizontal='left', vertical='center', wrap_text=False)
ws_rap.merge_cells('B4:L4')
ws_rap['B4'].alignment = Alignment(horizontal='left', vertical='center', wrap_text=False)

# KILAVUZ
ws_k = wb['KILAVUZ']
for mr in list(ws_k.merged_cells.ranges):
    if 'B4' in mr or 'B5' in mr:
        ws_k.merged_cells.remove(mr)
ws_k.merge_cells('B4:R4')
ws_k['B4'].alignment = Alignment(horizontal='left', vertical='center', wrap_text=False)
ws_k.merge_cells('B5:R5')
ws_k['B5'].alignment = Alignment(horizontal='left', vertical='center', wrap_text=False)

# 00_HAKKINDA_VE_REHBER
ws_info = wb['00_HAKKINDA_VE_REHBER']
for mr in list(ws_info.merged_cells.ranges):
    if 'B2' in mr or 'B3' in mr:
        ws_info.merged_cells.remove(mr)
ws_info.merge_cells('B2:F2')
ws_info['B2'].alignment = Alignment(horizontal='left', vertical='center', wrap_text=False)
ws_info.merge_cells('B3:F3')
ws_info['B3'].alignment = Alignment(horizontal='left', vertical='center', wrap_text=False)

wb.save(wb_path)
print("Tüm sayfalardaki başlık ve açıklama kırılmaları (wrap-text) başarıyla sıfırlandı.")
