import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import glob, sys, os

def standardize_workbook(filepath):
    try:
        wb = openpyxl.load_workbook(filepath)
    except Exception as e:
        return False, f"Yükleme hatası: {e}"

    modified = False

    # 1. TÜM SAYFALARDA SIFIR METİN KIRILMASI (A1 ve A3) & SÜTUN TAMPONU
    navy_fill = PatternFill(start_color='1B365D', end_color='1B365D', fill_type='solid')
    white_bold = Font(name='Montserrat', size=13, bold=True, color='FFFFFF')
    sub_font = Font(name='Segoe UI', size=9.5, italic=True, color='475569')

    for sname in wb.sheetnames:
        ws = wb[sname]
        ws.views.sheetView[0].showGridLines = False

        # Tablo genişliğini bul
        table_cols = [c for c in range(1, 30) if ws.cell(5, c).value is not None or ws.cell(4, c).value is not None]
        max_c = max(table_cols) if table_cols else 10
        if max_c < 8:
            max_c = 9
        col_let = get_column_letter(max_c)

        # A1 Banner Merge & Style
        if ws['A1'].value and isinstance(ws['A1'].value, str) and len(ws['A1'].value) > 10:
            # Eski merge temizle
            for mr in list(ws.merged_cells.ranges):
                if 'A1' in mr:
                    ws.merged_cells.remove(mr)
            ws.merge_cells(f'A1:{col_let}1')
            ws['A1'].font = white_bold
            ws['A1'].alignment = Alignment(horizontal='left', vertical='center', wrap_text=False)
            ws.row_dimensions[1].height = 32.0
            for c in range(1, max_c + 1):
                ws.cell(1, c).fill = navy_fill
            modified = True

        # A3 Alt Başlık Merge & Style (Sıfır Metin Kırılması)
        if ws['A3'].value and isinstance(ws['A3'].value, str) and len(ws['A3'].value) > 15:
            for mr in list(ws.merged_cells.ranges):
                if 'A3' in mr:
                    ws.merged_cells.remove(mr)
            ws.merge_cells(f'A3:{col_let}3')
            ws['A3'].font = sub_font
            ws['A3'].alignment = Alignment(horizontal='left', vertical='center', wrap_text=False)
            ws.row_dimensions[3].height = 22.0
            modified = True

        # Sütun tamponu (Rakam ve metin kesilmesini sıfırlama)
        for c in range(1, min(max_c + 2, 25)):
            c_letter = get_column_letter(c)
            cur_w = ws.column_dimensions[c_letter].width or 13.0
            max_l = 0
            for r in range(5, min(ws.max_row + 1, 40)):
                v = ws.cell(r, c).value
                if v is not None and not str(v).startswith('='):
                    max_l = max(max_l, len(str(v)))
            hdr = ws.cell(5, c).value
            if hdr:
                max_l = max(max_l, len(str(hdr)))
                needed_w = max(cur_w, max_l + 4.5)
                if needed_w > cur_w:
                    ws.column_dimensions[c_letter].width = round(needed_w, 1)
                    modified = True

    # 2. GAYRİRESMİ İFADELERİN TEMİZLENMESİ (MANDATE RULE 254)
    for sname in wb.sheetnames:
        ws = wb[sname]
        for row in ws.iter_rows(values_only=False):
            for cell in row:
                if cell.value and isinstance(cell.value, str):
                    val = cell.value
                    if any(k in val.lower() for k in ['5 yaş', 'çocuk', 'eli5']):
                        new_val = val.replace("5 yaşındaki bir çocuğun dahi anlayacağı", "Her kademedeki yöneticinin anında kavrayabileceği")
                        new_val = new_val.replace("5 Yaşındaki Birine Anlatır Gibi", "Yalın Kurumsal Dille")
                        new_val = new_val.replace("5 Yaşındaki Çocuğa Anlatır Gibi", "Yalın Kurumsal Dille")
                        new_val = new_val.replace("(5 YAŞINDA ÇOCUĞA ANLATIR GİBİ)", "")
                        cell.value = new_val.strip()
                        modified = True

    # 3. PANO/DASHBOARD GRAFİKLERİNİN DÜZENLENMESİ (MANDATE RULE 255 & 256)
    for p_name in ['PANO', 'DASHBOARD']:
        if p_name in wb.sheetnames:
            ws_p = wb[p_name]
            if ws_p._charts:
                for ch in ws_p._charts:
                    # Lejantı sağa sabitle
                    if hasattr(ch, 'legend') and ch.legend:
                        ch.legend.legendPos = "r"
                    # Halka grafiklerde dilim içi metin çakışmasını kapat
                    if type(ch).__name__ == 'DoughnutChart':
                        ch.dataLabels = None
                        ch.holeSize = 60
                    # Bar grafik
                    if type(ch).__name__ == 'BarChart':
                        if hasattr(ch, 'y_axis') and ch.y_axis:
                            ch.y_axis.title = "Tutar / Endeks"
                modified = True

    if modified:
        try:
            wb.save(filepath)
            return True, "Standardize edildi."
        except Exception as e:
            return False, f"Kaydetme hatası: {e}"
    else:
        return True, "Zaten tam uyumlu."

if __name__ == '__main__':
    files = glob.glob('EXCELARSIV_PRODUCTS/**/*.xlsx', recursive=True)
    files = [f for f in files if not f.split('/')[-1].startswith('~$')]
    
    print(f"Başlatılıyor: Toplam {len(files)} adet ürün anayasal standartlara geçiriliyor...")
    success_count = 0
    err_count = 0
    
    for idx, f in enumerate(files, start=1):
        ok, msg = standardize_workbook(f)
        if ok:
            success_count += 1
        else:
            err_count += 1
            print(f"[{idx}/{len(files)}] HATA: {f} -> {msg}")
        if idx % 20 == 0 or idx == len(files):
            print(f"İlerleme: {idx}/{len(files)} tamamlandı... (Başarılı: {success_count}, Hata: {err_count})")
            
    print(f"\nSONUÇ: {success_count} adet dosya başarıyla standardize edildi. Hatalı: {err_count}")
