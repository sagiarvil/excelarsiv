# calisma_prorokolu/EXCELARSIV_DECISION_OS_V15_1_RUNTIME_KERNEL.py içine Rule 256 ekle
with open('calisma_prorokolu/EXCELARSIV_DECISION_OS_V15_1_RUNTIME_KERNEL.py', 'r', encoding='utf-8') as f:
    content = f.read()

rule_256 = '''
    # MANDATE RULE 256: ZERO LABEL-DATA OVERLAP & ISOLATED CHART ANCHORING
    # 1. Grafik Başlık ve Lejant Çakışma Yasağı: Grafik lejantları (göstergeler) asla başlıkla veya veri çubuklarıyla
    #    aynı dikey eksene (üst/orta) konulamaz; daima sağ tarafa (legendPos='r') kilitlenerek başlık-lejant sürtünmesi engellenmelidir.
    # 2. Halka (Doughnut) Veri Etiketi Yığılması: Halka grafiklerde dilimlerin içine/üstüne uzun metinler (Kategori + Tutar)
    #    basılamaz (dataLabels=None); dağılım sağdaki temiz lejant ile okunmalıdır.
    # 3. Kılavuz Süreç Tablosu Sütun Genişlikleri: Kılavuz sayfalarındaki süreç adımları tablosunda Aşama Adı sütunu en az 30
    #    karakter genişliğinde olmalı, metinlerin 'AŞAMA', 'SİSTEM KURU' gibi yarım kesilmesine asla izin verilmemelidir.
'''

if 'MANDATE RULE 256' not in content:
    target = '    # MANDATE RULE 255: ZERO CHART-TEXT OVERLAP & COLUMN-COLLISION FORTRESS'
    content = content.replace(target, rule_256 + '\n' + target)
    with open('calisma_prorokolu/EXCELARSIV_DECISION_OS_V15_1_RUNTIME_KERNEL.py', 'w', encoding='utf-8') as f:
        f.write(content)
    print("Rule 256 Kernel py dosyasına başarıyla işlendi.")

# Master template senkronizasyonu
import openpyxl

master_path = 'calisma_prorokolu/EXCELARSIV_DECISION_OS_V15_1_MASTER_STANDART_UI_LOCKED.xlsx'
wb = openpyxl.load_workbook(master_path)

if '03_GORSEL_TASARIM' in wb.sheetnames:
    ws_v = wb['03_GORSEL_TASARIM']
    next_r = ws_v.max_row + 1
    ws_v[f'B{next_r}'] = "Sıfır Başlık-Lejant Çakışması & Kılavuz Tablo Genişlik Standardı"
    ws_v[f'C{next_r}'] = "Grafik lejantları başlığı ezmemesi için daima sağda (r) yer almalıdır. Halka grafiklerde dilim içi metin yığılması yasaktır. Kılavuz süreç tablolarında kolonlar metinleri 'SİSTEM KURU' gibi kesmeyecek biçimde en az 30 karakter genişlikte olmalıdır."
    ws_v[f'D{next_r}'] = "ZORUNLU"
    wb.save(master_path)
    print("Master Template: Rule 256 başarıyla işlendi.")
