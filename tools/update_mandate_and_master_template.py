import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# 1. calisma_prorokolu/EXCELARSIV_DECISION_OS_V15_1_MASTER_STANDART_UI_LOCKED.xlsx içine 00_HAKKINDA_VE_REHBER ekleme
master_path = 'calisma_prorokolu/EXCELARSIV_DECISION_OS_V15_1_MASTER_STANDART_UI_LOCKED.xlsx'
wb_m = openpyxl.load_workbook(master_path)

if '00_HAKKINDA_VE_REHBER' in wb_m.sheetnames:
    del wb_m['00_HAKKINDA_VE_REHBER']

ws_info = wb_m.create_sheet(title='00_HAKKINDA_VE_REHBER', index=1)
ws_info.views.sheetView[0].showGridLines = False
ws_info.sheet_properties.tabColor = '002563EB'

ws_info.column_dimensions['A'].width = 3.5
ws_info.column_dimensions['B'].width = 24.0
ws_info.column_dimensions['C'].width = 32.0
ws_info.column_dimensions['D'].width = 40.0
ws_info.column_dimensions['E'].width = 24.0
ws_info.column_dimensions['F'].width = 18.0

ws_info.row_dimensions[2].height = 36.0
ws_info['B2'] = "SİSTEM HAKKINDA, KULLANIM REHBERİ VE TERİMLER SÖZLÜĞÜ (SIEMENS STANDARDI)"
ws_info['B2'].font = Font(name='Montserrat', size=18, bold=True, color='0F172A')

ws_info.row_dimensions[3].height = 20.0
ws_info['B3'] = "5 Yaşındaki Çocuğa Anlatır Gibi Yalın Finansal Karar Mimarisi & Kurumsal Süreç Sözlüğü"
ws_info['B3'].font = Font(name='Segoe UI', size=10, italic=True, color='475569')

navy_fill = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
white_bold = Font(name='Segoe UI', size=9, bold=True, color='FFFFFF')
thin_gray = Side(style='thin', color='E2E8F0')

# Soru tablosu
ws_info['B5'] = "1. TEMEL MANTIK — 3 DAKİKADA SİSTEMİN ÖZÜ"
ws_info['B5'].font = Font(name='Montserrat', size=12, bold=True, color='1E3A8A')

questions = [
    ("SORU", "GÜNLÜK DİLLE CEVAP (5 YAŞINDA ÇOCUĞA ANLATIR GİBİ)", "ŞİRKETİNİZE FAYDASI NE?", "İLGİLİ SAYFA"),
    ("1. Bu dosya ne işe yarar?", 
     "Tıpkı bir uçağın kokpiti gibidir. Şirketin tüm borçlarını, taksitlerini ve kasadaki parasını tek bakışta gösterir.", 
     "Gözden kaçan taksitleri sıfırlar; faiz cezası ve itibar kaybını önler.", 
     "PANO, YÖNETİM_RAPORU"),
    ("2. Kullanıcı olarak ben ne yapacağım?", 
     "Yalnızca 'GİRDİLER' sekmesindeki sarı renkli hücrelere kredilerinizi yazacaksınız. Geri kalan her şeyi robot hesaplar.", 
     "Veri girişi 5 dakikada biter; karmaşık formüllerle veya kodlarla uğraşmazsınız.", 
     "GİRDİLER, KREDILER"),
    ("3. Sistem bana ne kazandırır?", 
     "Hangi bankanın pahalı olduğunu tespit eder. Borcu ucuza taşıyarak şirketinize somut nakit kazandırır.", 
     "Gereksiz faiz ve komisyonları yok eder; banka pazarlıklarında şirketinizi güçlü kılar.", 
     "REFINANSMAN, KARAR_MOTORU")
]

ws_info.row_dimensions[7].height = 24.0
for c_idx, title in enumerate(questions[0], start=2):
    c_let = openpyxl.utils.get_column_letter(c_idx)
    cell = ws_info[f'{c_let}7']
    cell.value = title
    cell.fill = navy_fill
    cell.font = white_bold
    cell.alignment = Alignment(horizontal='center' if c_let in ['B', 'E'] else 'left', vertical='center')

for idx, q in enumerate(questions[1:], start=8):
    ws_info.row_dimensions[idx].height = 42.0
    ws_info[f'B{idx}'] = q[0]
    ws_info[f'C{idx}'] = q[1]
    ws_info[f'D{idx}'] = q[2]
    ws_info[f'E{idx}'] = q[3]
    bg_color = 'F8FAFC' if idx % 2 == 0 else 'FFFFFF'
    row_fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type='solid')
    for col_let in ['B', 'C', 'D', 'E']:
        c = ws_info[f'{col_let}{idx}']
        c.fill = row_fill
        c.border = Border(top=thin_gray, bottom=thin_gray, left=thin_gray, right=thin_gray)
        c.alignment = Alignment(horizontal='center' if col_let in ['B', 'E'] else 'left', vertical='center', wrap_text=True)

# 03_GORSEL_TASARIM güncellemesi: Siemens Standardı Kuralı
if '03_GORSEL_TASARIM' in wb_m.sheetnames:
    ws_v = wb_m['03_GORSEL_TASARIM']
    next_r = ws_v.max_row + 1
    ws_v[f'B{next_r}'] = "Siemens & Enterprise UI Standartı"
    ws_v[f'C{next_r}'] = "00_HAKKINDA_VE_REHBER (5 yaşındaki çocuğa anlatır gibi sistem sözlüğü), kart içi sıfır gridline/artık border, simetrik dış çerçeve ve altın oran sütun genişlikleri zorunludur."
    ws_v[f'D{next_r}'] = "ZORUNLU"

wb_m.save(master_path)
print("Master UI Template Başarıyla Siemens Standartlarına Güncellendi.")
