# 1. calisma_prorokolu/EXCELARSIV_DECISION_OS_V15_1_RUNTIME_KERNEL.py içine Rule 255 ekle
with open('calisma_prorokolu/EXCELARSIV_DECISION_OS_V15_1_RUNTIME_KERNEL.py', 'r', encoding='utf-8') as f:
    content = f.read()

rule_255 = '''
    # MANDATE RULE 255: ZERO CHART-TEXT OVERLAP & COLUMN-COLLISION FORTRESS
    # 1. Grafik Tuval Alanı İzolasyonu: Grafiklerin altına denk gelen satırların (28..48 vb.) hücreleri tamamen boş
    #    olmalı, hiçbir metin veya tablo hücresi grafik nesnesinin arkasında ezilmemeli/baskılanmamalıdır.
    # 2. Grafik İçi Tipografi ve Lejant Yerleşimi: Göstergeler (Legends) ve Veri Etiketleri (DataLabels) asla grafik
    #    çizgilerinin veya pasta dilimlerinin üzerine taşmamalı; Bar grafiklerde tepeye (legendPos='t'), Halka
    #    grafiklerde sağa (legendPos='r') kilitlenerek etiket sürtünmesi sıfırlanmalıdır.
    # 3. Sayısal Kesilme Önleme (Column Auto-Padding): Veri tablolarında sayısal veya metinsel kolonlar asla dar
    #    bırakılamaz; sütun genişliği en az max(metin_uzunluğu) + 4.5 karakter tamponu içermeli ve asla '###' hatası üretmemelidir.
'''

if 'MANDATE RULE 255' not in content:
    target = '    # MANDATE RULE 254: C-LEVEL CORPORATE TONE & ZERO INFORMAL PHRASING'
    content = content.replace(target, rule_255 + '\n' + target)
    with open('calisma_prorokolu/EXCELARSIV_DECISION_OS_V15_1_RUNTIME_KERNEL.py', 'w', encoding='utf-8') as f:
        f.write(content)
    print("Rule 255 Kernel py dosyasına başarıyla işlendi.")
else:
    print("Rule 255 zaten mevcut.")

# 2. calisma_prorokolu/EXCELARSIV_DECISION_OS_V15_1_MASTER_STANDART_UI_LOCKED.xlsx içine kural ekle
import openpyxl

master_path = 'calisma_prorokolu/EXCELARSIV_DECISION_OS_V15_1_MASTER_STANDART_UI_LOCKED.xlsx'
wb = openpyxl.load_workbook(master_path)

if '03_GORSEL_TASARIM' in wb.sheetnames:
    ws_v = wb['03_GORSEL_TASARIM']
    next_r = ws_v.max_row + 1
    ws_v[f'B{next_r}'] = "Grafik Tuval İzolasyonu & Sıfır Ezilme/Sürtünme"
    ws_v[f'C{next_r}'] = "Grafik nesneleri asla arkasındaki hücreleri ezemez; altına gelen satırlar boş tuval olarak ayrılmalıdır. Grafik içi etiket ve lejantlar dilimlerle çakışmayacak konumlara (Bar: Üst, Halka: Sağ) yerleştirilmeli; tablolardaki sütunlar '###' veya rakam kesilmesine karşı en az +4.5 tampon paya sahip olmalıdır."
    ws_v[f'D{next_r}'] = "ZORUNLU"
    wb.save(master_path)
    print("Master Template: Grafik Tuval İzolasyonu kuralı başarıyla işlendi.")
