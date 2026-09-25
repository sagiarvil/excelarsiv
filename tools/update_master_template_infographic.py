import openpyxl

master_path = 'calisma_prorokolu/EXCELARSIV_DECISION_OS_V15_1_MASTER_STANDART_UI_LOCKED.xlsx'
wb = openpyxl.load_workbook(master_path)

if '03_GORSEL_TASARIM' in wb.sheetnames:
    ws_v = wb['03_GORSEL_TASARIM']
    
    # 1. Daimi Otomatik Senkronizasyon Kuralı
    next_r = ws_v.max_row + 1
    ws_v[f'B{next_r}'] = "Daimi Otomatik Protokol Senkronizasyonu (SSOT)"
    ws_v[f'C{next_r}'] = "Projedeki herhangi bir çalışmada geliştirilen her yeni tasarım standardı, mimari kural, formül motoru veya görsel iyileştirme; kullanıcı talep etmese dahi derhal bu master şablona ve runtime_kernel.py dosyasına işlenmek zorundadır. Bu iki dosya tek ve mutlak anayasadır."
    ws_v[f'D{next_r}'] = "ANAYASAL ZORUNLULUK"
    
    # 2. İnfografik Renk ve Vektör Grafik Standardı
    next_r += 1
    ws_v[f'B{next_r}'] = "İnfografik Renk Uyumu & Vektör Grafik Standardı"
    ws_v[f'C{next_r}'] = "PANO (Dashboard) ve Rapor sayfalarındaki grafikler ham Excel renkleriyle bırakılamaz. Uluslararası vektörel infografik renk paleti zorunludur: Birincil Hacim=#1B365D (Lacivert), İkincil Maliyet=#D97706 (Amber), Üçüncül Hedef=#059669 (Zümrüt Yeşil), Dağılım=#2563EB (Kraliyet Mavisi). Halka grafiklerde her dilim (DataPoint) bu paletle boyanmalı, delik oranı %60 kilitlenmelidir."
    ws_v[f'D{next_r}'] = "ZORUNLU"

wb.save(master_path)
print("Master Template: Auto-Sync Mandate ve İnfografik Renk Standardı başarıyla işlendi.")
